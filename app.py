import streamlit as st
from pypdf import PdfReader
from docx import Document

from langchain_google_genai import (
    ChatGoogleGenerativeAI,
    GoogleGenerativeAIEmbeddings
)

from langchain_core.documents import Document as LCDocument
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain.tools import tool
from langchain.agents import create_agent


# ==================================================
# PAGE CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="AI Interview Preparation Agent",
    page_icon="🎤",
    layout="wide"
)


# ==================================================
# TITLE
# ==================================================

st.title("🎤 AI Interview Preparation Agent")

st.write(
    "Upload your resume and practice a personalized "
    "AI-powered mock interview."
)


# ==================================================
# API KEY
# ==================================================

try:

    GOOGLE_API_KEY = st.secrets["GEMINI_API_KEY"]

except Exception:

    st.error("Gemini API key is not configured.")
    st.stop()


# ==================================================
# INITIALIZE GEMINI
# ==================================================

@st.cache_resource
def initialize_models():

    llm = ChatGoogleGenerativeAI(
        model="gemini-3.5-flash-lite",
        google_api_key=GOOGLE_API_KEY,
        temperature=0
    )

    embeddings = GoogleGenerativeAIEmbeddings(
        model="gemini-embedding-001",
        google_api_key=GOOGLE_API_KEY
    )

    return llm, embeddings


llm, embeddings = initialize_models()


# ==================================================
# SESSION STATE
# ==================================================

if "vector_store" not in st.session_state:
    st.session_state.vector_store = None

if "resume_name" not in st.session_state:
    st.session_state.resume_name = None

if "interview_agent" not in st.session_state:
    st.session_state.interview_agent = None

if "question" not in st.session_state:
    st.session_state.question = None

if "interview_started" not in st.session_state:
    st.session_state.interview_started = False

if "last_evaluation" not in st.session_state:
    st.session_state.last_evaluation = None


# ==================================================
# EXTRACT RESUME TEXT
# ==================================================

def extract_resume_text(uploaded_file):

    filename = uploaded_file.name.lower()

    # -------------------------------
    # PDF
    # -------------------------------

    if filename.endswith(".pdf"):

        reader = PdfReader(uploaded_file)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

        return text

    # -------------------------------
    # DOCX
    # -------------------------------

    elif filename.endswith(".docx"):

        document = Document(uploaded_file)

        text = ""

        for paragraph in document.paragraphs:

            text += paragraph.text + "\n"

        return text

    # -------------------------------
    # Unsupported file
    # -------------------------------

    else:

        raise ValueError(
            "Only PDF and DOCX files are supported."
        )


# ==================================================
# RESUME UPLOAD
# ==================================================

uploaded_file = st.file_uploader(
    "Upload your Resume",
    type=["pdf", "docx"]
)


# ==================================================
# PROCESS RESUME
# ==================================================

if uploaded_file is not None:

    if (
        st.session_state.resume_name
        != uploaded_file.name
    ):

        with st.spinner("Processing resume..."):

            try:

                # ----------------------------------
                # Extract resume text
                # ----------------------------------

                resume_text = extract_resume_text(
                    uploaded_file
                )

                if not resume_text.strip():

                    st.error(
                        "Could not extract text from the resume."
                    )

                    st.stop()

                # ----------------------------------
                # Split resume
                # ----------------------------------

                text_splitter = RecursiveCharacterTextSplitter(
                    chunk_size=1000,
                    chunk_overlap=200
                )

                resume_chunks = text_splitter.split_text(
                    resume_text
                )

                # ----------------------------------
                # Create LangChain documents
                # ----------------------------------

                documents = [
                    LCDocument(
                        page_content=chunk,
                        metadata={
                            "source": uploaded_file.name
                        }
                    )
                    for chunk in resume_chunks
                ]

                # ----------------------------------
                # Create FAISS vector database
                # ----------------------------------

                vector_store = FAISS.from_documents(
                    documents,
                    embeddings
                )

                st.session_state.vector_store = (
                    vector_store
                )

                st.session_state.resume_name = (
                    uploaded_file.name
                )

                # ----------------------------------
                # Resume retrieval tool
                # ----------------------------------

                @tool
                def retrieve_resume_context(
                    query: str
                ):
                    """
                    Retrieve relevant information
                    from the candidate's resume.
                    """

                    docs = vector_store.similarity_search(
                        query,
                        k=3
                    )

                    return "\n\n".join(
                        f"Content: {doc.page_content}"
                        for doc in docs
                    )

                # ----------------------------------
                # Agent system prompt
                # ----------------------------------

                system_prompt = """
You are an AI Interview Preparation Agent.

Your job is to conduct a personalized mock
interview using the candidate's resume.

Your responsibilities:

1. Ask technical interview questions.
2. Ask questions based on the candidate's projects.
3. Ask questions based on technical skills.
4. Ask questions about education and experience.
5. Evaluate candidate answers.
6. Give a score out of 10.
7. Explain what was correct.
8. Explain what was missing.
9. Suggest how the answer can be improved.
10. Provide a better sample answer.
11. Ask a relevant follow-up question.
12. Never invent information that is not present
    in the resume.

Keep the interview suitable for a college student
preparing for placements.

Use the resume retrieval tool whenever
resume-specific information is needed.
"""

                # ----------------------------------
                # Create interview agent
                # ----------------------------------

                interview_agent = create_agent(
                    llm,
                    [retrieve_resume_context],
                    system_prompt=system_prompt
                )

                st.session_state.interview_agent = (
                    interview_agent
                )

                # ----------------------------------
                # Reset interview
                # ----------------------------------

                st.session_state.question = None

                st.session_state.interview_started = (
                    False
                )

                st.session_state.last_evaluation = (
                    None
                )

                st.success(
                    "✅ Resume processed successfully!"
                )

            except Exception as e:

                st.error(
                    f"Error processing resume: {str(e)}"
                )


# ==================================================
# INTERVIEW CONFIGURATION
# ==================================================

if st.session_state.vector_store is not None:

    st.markdown("---")

    st.header("⚙️ Interview Configuration")

    job_role = st.text_input(
        "Target Job Role",
        placeholder="Example: Java Developer"
    )

    interview_type = st.selectbox(
        "Interview Type",
        [
            "Technical Interview",
            "HR Interview",
            "Project-Based Interview",
            "Mixed Interview"
        ]
    )


# ==================================================
# START INTERVIEW
# ==================================================

if (
    st.session_state.vector_store is not None
    and job_role
):

    if st.button(
        "🚀 Start Interview",
        type="primary"
    ):

        with st.spinner(
            "Preparing your first interview question..."
        ):

            start_prompt = f"""
Start a mock interview for the candidate.

Target job role:
{job_role}

Interview type:
{interview_type}

Ask the FIRST interview question.

The question should be personalized using
information from the candidate's resume.

Return only the interview question.
"""

            try:

                result = (
                    st.session_state.interview_agent.invoke(
                        {
                            "messages": [
                                {
                                    "role": "user",
                                    "content": start_prompt
                                }
                            ]
                        }
                    )
                )

                # ----------------------------------
                # Extract question text
                # ----------------------------------

                content = result[
                    "messages"
                ][-1].content

                if isinstance(content, list):

                    question = "".join(
                        item.get("text", "")
                        for item in content
                        if isinstance(item, dict)
                    )

                else:

                    question = str(content)

                # ----------------------------------
                # Save question
                # ----------------------------------

                st.session_state.question = (
                    question
                )

                st.session_state.interview_started = (
                    True
                )

                st.session_state.last_evaluation = (
                    None
                )

            except Exception as e:

                st.error(
                    f"Error starting interview: {str(e)}"
                )


# ==================================================
# DISPLAY INTERVIEW QUESTION
# ==================================================

if (
    st.session_state.interview_started
    and st.session_state.question
):

    st.markdown("---")

    st.header("🎯 Interview Question")

    st.info(
        st.session_state.question
    )

    # ----------------------------------------------
    # Candidate answer
    # ----------------------------------------------

    answer = st.text_area(
        "Your Answer",
        height=180,
        placeholder="Type your interview answer here..."
    )

    # ----------------------------------------------
    # Submit answer
    # ----------------------------------------------

    if st.button("📊 Submit Answer"):

        if not answer.strip():

            st.warning(
                "Please enter your answer first."
            )

        else:

            with st.spinner(
                "Evaluating your answer..."
            ):

                evaluation_prompt = f"""
Evaluate the candidate's interview answer.

Target job role:
{job_role}

Interview type:
{interview_type}

Interview question:
{st.session_state.question}

Candidate answer:
{answer}

Provide the evaluation in exactly this structure:

## Score

Give a score out of 10.

## What You Did Well

Explain the strong points.

## What Was Missing

Explain what could be improved.

## Improvement Suggestions

Give practical suggestions.

## Better Sample Answer

Provide a stronger sample answer.

## Follow-up Question

Ask one relevant follow-up interview question.

Use the candidate's resume when relevant.

Do not invent resume information.
"""

                try:

                    result = (
                        st.session_state.interview_agent.invoke(
                            {
                                "messages": [
                                    {
                                        "role": "user",
                                        "content": evaluation_prompt
                                    }
                                ]
                            }
                        )
                    )

                    # ----------------------------------
                    # Extract evaluation text
                    # ----------------------------------

                    content = result[
                        "messages"
                    ][-1].content

                    if isinstance(content, list):

                        evaluation = "".join(
                            item.get("text", "")
                            for item in content
                            if isinstance(item, dict)
                        )

                    else:

                        evaluation = str(content)

                    # ----------------------------------
                    # Save evaluation
                    # ----------------------------------

                    st.session_state.last_evaluation = (
                        evaluation
                    )

                except Exception as e:

                    st.error(
                        f"Error evaluating answer: {str(e)}"
                    )


# ==================================================
# DISPLAY EVALUATION
# ==================================================

if st.session_state.last_evaluation:

    st.markdown("---")

    st.header("📊 Answer Evaluation")

    st.markdown(
        st.session_state.last_evaluation
    )


# ==================================================
# FOOTER
# ==================================================

st.markdown("---")

st.caption(
    "AI Interview Preparation Agent | "
    "Gemini + LangChain + RAG + FAISS"
)
